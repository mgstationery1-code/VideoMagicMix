from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import ScriptScheme, ScriptSegment

ROW_RE = re.compile(r"^(Hook|Body\d+)$", re.IGNORECASE)


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value or "")
    value = value.replace("\u3000", "")
    return re.sub(r"[\s\W_]+", "", value, flags=re.UNICODE).lower()


def _col_letter(index: int) -> str:
    out = ""
    while index:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def import_workbook(path: str) -> tuple[str, list[ScriptScheme], list[str], list[str]]:
    """Import the workbook while refusing visibly corrupted replacement characters."""
    file_path = Path(path)
    if not file_path.is_file() or file_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Excel 文件必须是 .xlsx 或 .xlsm")
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取 Excel：{exc}") from exc
    sheet = workbook.active
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    errors: list[str] = []
    warnings: list[str] = []
    for r, row in enumerate(values, 1):
        is_hook_row = bool(row) and _text(row[0]).lower() == "hook"
        for c, value in enumerate(row, 1):
            if is_hook_row and c >= 3:
                continue
            text = _text(value)
            if "\ufffd" in text:
                errors.append(f"{sheet.title}!{_col_letter(c)}{r} 含有替换字符，疑似中文编码损坏")
    data_row_indices = [i for i, row in enumerate(values) if ROW_RE.match(_text(row[0] if row else ""))]
    if not data_row_indices:
        errors.append("未找到 Hook 或 Body1、Body2 等数据行")
        return sheet.title, [], errors, warnings
    first_data = data_row_indices[0]
    header_index = max(0, first_data - 1)
    header = values[header_index]
    scheme_columns = [i for i in range(2, len(header)) if _text(header[i])]
    if not scheme_columns:
        errors.append("未找到 C 列及之后的字幕方案列")
        return sheet.title, [], errors, warnings
    seen_keys: set[str] = set()
    schemes: list[ScriptScheme] = []
    for ordinal, col in enumerate(scheme_columns, 1):
        display_name = _text(header[col]) or f"字幕{ordinal}"
        scheme = ScriptScheme(scheme_id=f"scheme-{ordinal}", source_column=_col_letter(col + 1), display_name=display_name)
        for row_index in data_row_indices:
            row = values[row_index]
            row_key = _text(row[0] if len(row) > 0 else "")
            group_name = _text(row[1] if len(row) > 1 else "")
            text = _text(row[col] if col < len(row) else "")
            is_hook = row_key.lower() == "hook"
            # Hook is a finished clip with its own voice and subtitles.  Keep the
            # row only for workbook fidelity; it is never part of Body narration.
            enabled = bool(text) and not is_hook
            skip_reason = None
            if not text:
                if not is_hook:
                    skip_reason = "文案为空，已跳过"
                    warnings.append(f"{display_name} / {row_key}：{skip_reason}")
            elif is_hook:
                skip_reason = "成品 Hook 自带配音和字幕，Excel Hook 文案已忽略"
            if row_key.lower() in seen_keys and ordinal == 1:
                errors.append(f"{sheet.title}!A{row_index + 1} 重复的行标识：{row_key}")
            if ordinal == 1:
                seen_keys.add(row_key.lower())
            scheme.segments.append(ScriptSegment(row_key=row_key, group_name=group_name, text=text, normalized_text=normalize_text(text), enabled=enabled, skip_reason=skip_reason))
        schemes.append(scheme)
    workbook.close()
    return sheet.title, schemes, errors, warnings


def preview_scheme(scheme: ScriptScheme, length: int = 10) -> list[dict[str, Any]]:
    return [{"row_key": segment.row_key, "group_name": segment.group_name, "preview": segment.text[:length], "enabled": segment.enabled, "skip_reason": segment.skip_reason} for segment in scheme.segments]
